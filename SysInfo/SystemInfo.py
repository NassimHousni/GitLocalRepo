from pathlib import Path
from openpyxl import Workbook, load_workbook

from Filteren.OefLes3_Filteren import collect_system_info


# Gebruik de bestaande functie uit oefening 3.
if __name__ == "__main__":
    info = collect_system_info()
       # Selecteer de gegevens voor de kolommen in Excel.
    cpu_freq = info["cpu"]["cpu_freq"]

    filtered_info = {
        "computernaam": info["network"]["hostname"],
        "processor": info["cpu"]["processor"],
        "physical_cores": info["cpu"]["physical_cores"],
        "max_cpu_frequency_mhz": cpu_freq["max"] if cpu_freq else None,
        "total_ram_gib": round(info["memory"]["total"] / (1024 ** 3), 2),
    }

    print(filtered_info)

        # Open het bestaande Excel-bestand of maak een nieuw werkboek.
    excel_file = Path(__file__).resolve().parent / "system_info.xlsx"

    if excel_file.exists():
        wb = load_workbook(excel_file)
        ws = wb["System Info"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "System Info"
        ws.append(list(filtered_info.keys()))

            # Werk de bestaande computer bij of voeg een nieuwe rij toe.
    for row_number in range(2, ws.max_row + 1):
        if ws.cell(row=row_number, column=1).value == filtered_info["computernaam"]:
            for column_number, value in enumerate(filtered_info.values(), start=1):
                ws.cell(row=row_number, column=column_number, value=value)
            break
    else:
        ws.append(list(filtered_info.values()))

    # Sla de gegevens op in het Excel-bestand.
    wb.save(excel_file)
    print(f"Excel-bestand opgeslagen: {excel_file}")